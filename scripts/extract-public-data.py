"""Rebuild the site's public aggregate data from downloaded government files.
Usage: python scripts/extract-public-data.py /path/to/research
Reads only institutional mosque records and published aggregate census tables.
"""
import hashlib
import json
import re
import sys
from pathlib import Path
import openpyxl
import pdfplumber

src = Path(sys.argv[1])
dest = Path(__file__).resolve().parents[1] / 'public' / 'data'
dest.mkdir(exist_ok=True)

def names(value):
    value = value.strip()
    en = re.search(r'[A-Za-z][A-Za-z .&/\-\'()]*$', value).group().strip()
    si = ' '.join(re.findall(r'[\u0d80-\u0dff][\u0d80-\u0dff\u200c\u200d\s.]*', value)).strip()
    ta = ' '.join(re.findall(r'[\u0b80-\u0bff][\u0b80-\u0bff\u200c\u200d\s.]*', value)).strip()
    return {'en': en, 'ta': ' '.join(ta.split()), 'si': ' '.join(si.split())}

def write(name, value):
    (dest / name).write_text(json.dumps(value, ensure_ascii=False, separators=(',', ':')), encoding='utf8')

gnbook = openpyxl.load_workbook(src / 'gn-retry.xlsx', data_only=True, read_only=True)
gnrows = [r for r in gnbook.active.values if isinstance(r[0], int)]
assert len(gnrows) == 14003
assert sum(r[6] for r in gnrows) == 21781800
assert sum(r[9] for r in gnrows if isinstance(r[9], int)) == 2327605
codes = {r[1]: str(r[0]) for r in gnrows}
gn_ds = {(r[1], r[3]): str(r[2]).zfill(2) for r in gnrows}

districts = []
for row in openpyxl.load_workbook(src / 'district-religion.xlsx', data_only=True).active.values:
    if not isinstance(row[0], str) or not isinstance(row[1], int) or row[1] < 100:
        continue
    name = names(row[0])
    item = {'names': name, 'total': row[1], 'islam': row[4]}
    if name['en'] == 'Sri Lanka':
        national = item
    else:
        item['id'] = codes[name['en']]
        districts.append(item)
assert len(districts) == 25
assert sum(r['total'] for r in districts) == national['total']
assert sum(r['islam'] for r in districts) == national['islam']

# Explicit spelling aliases only; these do not assign mosques to census areas.
aliases = {
    ('Kandy', 'Delthota'): 'Deltota',
    ('Nuwara Eliya', 'Thalawakele'): 'Thalawakelle',
    ('Nuwara Eliya', 'Ambagamuwa'): 'Ambagamuwa Koralaya',
    ('Galle', 'Benthota'): 'Bentota',
    ('Matara', 'Weligam'): 'Weligama',
    ('Hambantota', 'Lunugamwehera'): 'Lunugamvehera',
    ('Jaffna', 'Valikamam North'): 'Valikamam North (Tellipallai)',
    ('Jaffna', 'Vadamaradchy South-West (Karaveddy)'): 'Vadamaradchi South-West (Karaveddy)',
    ('Jaffna', 'Vadamaradchy East'): 'Vadamaradchi East',
    ('Jaffna', 'Vadamaradchy North (Point Perdro)'): 'Vadamaradchi North (Point Pedro)',
    ('Jaffna', 'Thenmaradchy (Chavakachcheri)'): 'Thenmaradchi (Chavakachcheri)',
    ('Mullaitivu', 'Puthukudiyiruppu'): 'Puthukkudiyiruppu',
    ('Mullaitivu', 'Oddusudan'): 'Oddusuddan',
    ('Batticaloa', 'Poratheevu Pattu'): 'Porativu Pattu',
    ('Batticaloa', 'Manmunai South & Eruvil Pattu'): 'Manmunai South & Eruvil pattu',
    ('Ampara', 'Sainthamarathu'): 'Sainthamaruthu',
    ('Trincomalee', 'Padavisripura'): 'Padavi Sri Pura',
    ('Ratnapura', 'Kalthota'): 'Kaltota'
}
divisions = []
district = None
unmatched = []
for row in openpyxl.load_workbook(src / 'ds-religion.xlsx', data_only=True).active.values:
    if not isinstance(row[1], str) or not isinstance(row[2], int) or row[2] < 100:
        continue
    name = names(row[1])
    if name['en'] == 'Sri Lanka':
        continue
    if name['en'].endswith(' District'):
        district = name['en'].removesuffix(' District')
        continue
    key = (district, aliases.get((district, name['en']), name['en']))
    if key not in gn_ds:
        unmatched.append(key)
        continue
    divisions.append({'id': codes[district] + '-' + gn_ds[key], 'district': codes[district], 'names': name, 'total': row[2], 'islam': row[5]})
if unmatched:
    print('Unmatched DS names:', unmatched)
    print('Available candidates:', sorted(gn_ds))
    raise SystemExit(1)
assert len(divisions) == 340
for ds in divisions:
    members = [r for r in gnrows if f'{r[0]}-{r[2]:02}' == ds['id']]
    assert sum(r[6] for r in members) == ds['total'], ds
    assert sum(r[9] for r in members if isinstance(r[9], int)) <= ds['islam'], ds
for d in districts:
    subset = [s for s in divisions if s['district'] == d['id']]
    assert sum(s['total'] for s in subset) == d['total'], d
    assert sum(s['islam'] for s in subset) == d['islam'], d
    gns = [r for r in gnrows if str(r[0]) == d['id']]
    assert sum(r[6] for r in gns) == d['total'], d
    d['gnCount'] = len(gns)
    d['dsCount'] = len(subset)
    write('gn-' + d['id'] + '.json', [
        {'id': f'{r[0]}-{r[2]:02}-{r[4]:03}', 'ds': f'{r[0]}-{r[2]:02}',
         'name': r[5], 'total': r[6], 'islam': r[9] if isinstance(r[9], int) else None}
        for r in gns
    ])

mosques = []
for page_no, page in enumerate(pdfplumber.open(src / 'mosques-anuradhapura.pdf').pages, 1):
    for table in page.extract_tables():
        for r in table:
            if not r[0] or not r[0].isdigit():
                continue
            no, registration, name, typ, address, city = [' '.join((s or '').split()) for s in r]
            mosques.append({'entry': int(no), 'registration': registration, 'name': name, 'type': typ, 'address': address, 'city': city, 'page': page_no})
assert [m['entry'] for m in mosques] == list(range(1, 109))
assert mosques[96]['registration'] == 'R/2141/A 112'
assert mosques[96]['name'] == 'JAMIA KADHEEJA JUMMAH MOSQUE'
assert mosques[96]['address'] == '516/23, UDHAYA MAWATHA,2ND BUILDING'
write('mosques-anuradhapura.json', mosques)
write('population.json', {'year': 2024, 'national': national, 'districts': districts, 'divisions': divisions})
files = {
    'district-religion.xlsx': 'https://www.statistics.gov.lk/Population/StaticalInformation/CPH2024/PopulationA/A3',
    'ds-religion.xlsx': 'https://www.statistics.gov.lk/Population/StaticalInformation/CPH2024/PopulationA/A7',
    'gn-retry.xlsx': 'https://www.statistics.gov.lk/Population/StaticalInformation/CPH2024/GNLevel/GN_Level_Population_by_Religion',
    'mosques-anuradhapura.pdf': 'https://muslimaffairs.gov.lk/wp-content/uploads/2023/12/ANURADHAPURA-DISTRICT.pdf'
}
write('provenance.json', {'checked': '2026-10-04', 'censusYear': 2024,
    'sources': [{'file': n, 'url': url, 'sha256': hashlib.sha256((src / n).read_bytes()).hexdigest()} for n, url in files.items()],
    'notes': ['Religion category: Islam; persons, not households or ethnicity.',
    'GN religion counts under 10 are grouped into Other by DCS. Source dashes are null, not zero.',
    'GN Islam counts are not summed to replace district or DS totals.',
    'No mosque boundary, resident affiliation, or mosque-to-GN match is inferred.',
    'Mosque names, registration strings and city spellings are retained from the official list. Whitespace is normalised.']})
print('Validated: 108 mosque entries; 25 districts; 340 DS divisions; 14,003 GN divisions.')
print('Anuradhapura:', next(d for d in districts if d['names']['en'] == 'Anuradhapura'))
